"""Gera docs/Sprint3_Cybersecurity_Ford_VINShare.docx (documento único, separado por atividade).

Trechos de código são lidos dos arquivos reais do repositório; resultados de testes, ferramentas
e alertas vêm de docs/evidencias/ e observability/, então o documento sempre reflete o estado atual.
Rode depois dos testes, das ferramentas, de simular_trafego.py e de gerar_figuras.py.
"""
import json
import defusedxml.ElementTree as ET  # parser seguro contra XXE e "XML bomb" (Semgrep use-defused-xml-parse)
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "Sprint3_Cybersecurity_Ford_VINShare.docx"

FORD = RGBColor(0x13, 0x3A, 0x7C)
INK2 = RGBColor(0x52, 0x51, 0x4E)
HEADER_FILL = "133A7C"
ZEBRA_FILL = "F3F5F9"
CODE_FILL = "F4F4F1"

doc = Document()

# ------------------------------------------------------------------ estilos base
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
CONTENT_W = 17.0

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

for name, size, before in (("Heading 1", 17, 18), ("Heading 2", 13, 12), ("Heading 3", 11.5, 9)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = FORD
    rfonts = st.element.rPr.find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:ascii"), "Calibri")
        rfonts.set(qn("w:hAnsi"), "Calibri")
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True


def shade(el, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


# ------------------------------------------------------------------ helpers de conteúdo
def h1(text, page_break=True):
    heading = doc.add_heading(text, level=1)
    heading.paragraph_format.page_break_before = page_break
    heading.paragraph_format.space_before = Pt(0)
    return heading


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def p(text="", italic=False, size=None, color=None, align=None, bold=False, space_after=None):
    para = doc.add_paragraph()
    _rich(para, text, italic=italic, size=size, color=color, bold=bold)
    if align:
        para.alignment = align
    if space_after is not None:
        para.paragraph_format.space_after = Pt(space_after)
    return para


def _rich(para, text, italic=False, size=None, color=None, bold=False):
    """**negrito**, *itálico* e `código` inline."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = para.add_run(part[1:-1])
            run.italic = True
            run.bold = bold
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
            continue
        elif part.startswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = para.add_run(part)
            run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def bullet(text, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    _rich(para, text)
    para.paragraph_format.space_after = Pt(2)
    return para


def table(headers, rows, widths, font=8.5, header_font=8.5):
    assert abs(sum(widths) - CONTENT_W) < 0.05, (headers, sum(widths))
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    tr_pr.append(rep)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
        para = cell.paragraphs[0]
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(header_font)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Cm(widths[i])
            if r % 2 == 1:
                shade(cell._tc.get_or_add_tcPr(), ZEBRA_FILL)
            lines = str(text).split("\n")
            para = cell.paragraphs[0]
            for j, line in enumerate(lines):
                if j:
                    para = cell.add_paragraph()
                _rich(para, line, size=font)
                para.paragraph_format.space_after = Pt(0)
    # bordas discretas
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "C3C2B7")
        borders.append(b)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code(text, caption=None, size=7.8):
    if caption:
        cap = p(caption, italic=True, size=8.5, color=INK2, space_after=1)
        cap.paragraph_format.keep_with_next = True
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    shade(pPr, CODE_FILL)
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), HEADER_FILL)
    bdr.append(left)
    pPr.append(bdr)
    para.paragraph_format.left_indent = Cm(0.25)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.0
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(size)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if i < len(lines) - 1:
            run.add_break()
    return para


def snippet(rel_path, start, end=None, extra=0):
    """Trecho de um arquivo real: da 1ª linha que contém `start` até a 1ª seguinte com `end`
    (mais `extra` linhas). Sem `end`, só a linha inicial."""
    lines = (ROOT / rel_path).read_text(encoding="utf-8").splitlines()
    i = next(n for n, l in enumerate(lines) if start in l)
    j = i if end is None else next(n for n in range(i + 1, len(lines)) if end in lines[n])
    block = lines[i:j + 1 + extra]
    indent = min(len(l) - len(l.lstrip()) for l in block if l.strip())
    return "\n".join(l[indent:] for l in block)


def figure(path, caption, width=CONTENT_W):
    doc.add_picture(str(path), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p(caption, italic=True, size=8.5, color=INK2, align=WD_ALIGN_PARAGRAPH.CENTER)


def callout(text):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Cm(CONTENT_W)
    shade(cell._tc.get_or_add_tcPr(), "EEF4FC")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4" if edge != "left" else "24")
        b.set(qn("w:color"), HEADER_FILL if edge == "left" else "C9D8EE")
        borders.append(b)
    t._tbl.tblPr.append(borders)
    _rich(cell.paragraphs[0], text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ------------------------------------------------------------------ dados de evidência
import re
import sys

sys.path.insert(0, str(ROOT / "observability"))
from avaliar_alertas import evaluate, load_series as load_alert_series  # noqa: E402

EVID = DOCS / "evidencias"


def load_junit(name):
    root = ET.parse(EVID / name).getroot()
    suite = root if root.tag == "testsuite" else root.find("testsuite")
    cases = [(tc.get("classname").split(".")[-1], tc.get("name"),
              "FALHOU" if tc.find("failure") is not None or tc.find("error") is not None else "PASSOU")
             for tc in suite.iter("testcase")]
    return cases


api_cases = load_junit("pytest-report.xml")
iot_cases = load_junit("pytest-iot-report.xml")
all_cases = api_cases + iot_cases
n_tests = len(all_cases)
n_fail = sum(1 for *_, r in all_cases if r != "PASSOU")
by_class = Counter(c for c, _, _ in all_cases)

logs = [json.loads(line) for line in (ROOT / "observability" / "exemplos_logs.jsonl").read_text(encoding="utf-8").splitlines()]
event_counts = Counter(l.get("event") for l in logs)
alert_report = evaluate()
alert_ticks = load_alert_series()
old_bruteforce_peak = max(t["login_failure"] for t in alert_ticks)
old_bola_peak = max(t["authz_denied"] for t in alert_ticks)

semgrep = json.loads((EVID / "semgrep.json").read_text(encoding="utf-8"))
semgrep_files = len(semgrep["paths"]["scanned"])
semgrep_findings = len(semgrep["results"])
gitleaks_findings = len(json.loads((EVID / "gitleaks.json").read_text(encoding="utf-8") or "[]"))
gitleaks_bytes = re.search(r"scanned ~(\d+) bytes", (EVID / "gitleaks.txt").read_text(encoding="utf-8"))
gitleaks_kb = int(gitleaks_bytes.group(1)) // 1024 if gitleaks_bytes else 0
bandit_txt = (EVID / "bandit.txt").read_text(encoding="utf-8")
bandit_loc = re.search(r"Total lines of code: (\d+)", bandit_txt).group(1)
bandit_clean = "No issues identified" in bandit_txt
bandit_full = (EVID / "bandit-completo.txt").read_text(encoding="utf-8")
bandit_sim = len(re.findall(r"Location: observability", bandit_full))
pip_audit_clean = "No known vulnerabilities" in (EVID / "pip-audit.txt").read_text(encoding="utf-8")
rules_txt = (EVID / "semgrep-regras.txt").read_text(encoding="utf-8")


def first_log(event, **match):
    return next(l for l in logs if l.get("event") == event and all(l.get(k) == v for k, v in match.items()))


# ------------------------------------------------------------------ rodapé / cabeçalho
def add_field(run, instr):
    for tag, attrs in (("w:fldChar", {"w:fldCharType": "begin"}),
                       ("w:instrText", None),
                       ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if attrs:
            for k, v in attrs.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hr = hp.add_run("Challenge Ford 2026 · Sprint 3 · Cybersecurity · Ford VIN Share")
hr.font.size = Pt(8)
hr.font.color.rgb = INK2
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run()
fr.font.size = Pt(8)
add_field(fr, "PAGE")

# ================================================================== CAPA
for _ in range(5):
    p()
p("FIAP · Challenge Ford 2026", size=12, color=INK2, space_after=2)
p("Sprint 3 — Cybersecurity", size=26, bold=True, color=FORD, space_after=2)
p("DevSecOps aplicado à solução Ford VIN Share", size=15, color=FORD, space_after=14)
p("Desafio 02 — Impulsionando o VIN Share na América do Sul com Soluções Inteligentes",
  size=11, color=INK2, space_after=40)
table(["Item", "Informação"], [
    ["Grupo", "[preencher]"],
    ["Integrantes (nome — RM)", "[preencher]\n[preencher]\n[preencher]"],
    ["Turma", "[preencher]"],
    ["Repositório GitHub", "[preencher link]"],
    ["Entrega", "27/09/2026"],
], [5.0, 12.0], font=10, header_font=10)
p()
p("Documento único, organizado pelas quatro subetapas da Sprint: (1) Pipeline DevSecOps Integrado, "
  "(2) Segurança em Código e Infraestrutura, (3) Observabilidade, Monitoramento e Resposta e "
  "(4) Compliance, Riscos e Segurança Contínua. Cobre API, mobile, IoT, dados, ML e arquitetura.",
  size=9.5, color=INK2)

# ================================================================== SUMÁRIO
h1("Sumário")
for item in [
    "0. Contexto da solução, arquitetura e rastreabilidade dos requisitos",
    "1. Pipeline DevSecOps Integrado (peso 3,0)",
    "2. Segurança em Código e Infraestrutura (peso 2,5)",
    "3. Observabilidade, Monitoramento e Resposta (peso 2,0)",
    "4. Compliance, Riscos e Segurança Contínua (peso 2,5)",
    "Anexo A. Estrutura do repositório e como executar",
    "Anexo B. Lista completa dos testes automatizados",
]:
    p(item, size=11, space_after=4)

# ================================================================== 0. CONTEXTO
h1("0. Contexto da solução e arquitetura de segurança")
p("O **VIN Share** mede a porcentagem de veículos Ford que fazem manutenção na rede oficial. A solução do "
  "grupo ajuda as concessionárias a reter clientes no pós-venda com quatro partes: (a) um **app mobile** "
  "(Expo/React Native) em que consultores e gestores trabalham leads de serviço; (b) uma **API REST** "
  "(FastAPI) que expõe leads, indicadores de Service Share e predições; (c) **telemetria de veículos "
  "conectados** recebida por MQTT e validada por um serviço de ingestão; e (d) um **modelo de Machine "
  "Learning** que estima a probabilidade de o veículo sair da rede Ford (churn).")
p("A solução trata dados pessoais (nome, CPF, telefone), dados do veículo ligados ao proprietário (VIN, "
  "quilometragem, códigos de falha) e dados de desempenho das concessionárias. Cada controle descrito aqui "
  "existe no repositório e, sempre que possível, foi **executado e verificado** (testes, SAST, SCA, "
  "secret scanning, scan de IaC e simulação de ataques).")

figure(ROOT / "docs" / "img" / "arquitetura_seguranca.png",
       "Figura 1 — Arquitetura de segurança: zonas de confiança, controles por componente e fluxos protegidos.",
       width=14.0)

h2("0.1 Perfis de acesso")
p("O enunciado cita os perfis Brigadista, Gestor e Administrador. No contexto do VIN Share adaptamos os mesmos "
  "três níveis (operacional, gestão e administração) para a realidade das concessionárias:")
table(["Perfil", "Equivalente no enunciado", "Quem é", "Escopo de dados"], [
    ["consultor", "Brigadista (operacional)", "Consultor de serviço da concessionária",
     "Somente leads da própria concessionária"],
    ["gestor", "Gestor", "Gestor de pós-venda da concessionária",
     "Leads + Service Share + predições da própria concessionária"],
    ["admin", "Administrador", "Equipe Ford (analytics / segurança)",
     "Todas as concessionárias, trilha de auditoria e usuários"],
], [2.4, 3.8, 5.2, 5.6])

h2("0.2 Stack e escopo desta entrega")
table(["Camada", "Tecnologia", "O que foi entregue e verificado nesta Sprint"], [
    ["API", "Python 3.13 · FastAPI · PyJWT · bcrypt · cryptography",
     f"Hardening completo + {len(api_cases)} testes de segurança executados"],
    ["Mobile", "Expo / React Native · expo-secure-store",
     "Sessão segura, versão mínima obrigatória (426) e hardening do Android"],
    ["IoT", "MQTT (Mosquitto 2.0) · paho-mqtt",
     f"Broker TLS/mTLS/ACL/CRL, ingestão com anti-replay + {len(iot_cases)} testes (inclui handshake mTLS real)"],
    ["Dados", "PostgreSQL (Azure) · Key Vault", "PII cifrada por campo (AES-256-GCM), índice cego, rotina de backup"],
    ["ML", "Modelo de churn (regressão logística)",
     "Artefato JSON com SHA-256 obrigatório em produção, entradas validadas, métricas de drift"],
    ["Infra", "Docker · Kubernetes (AKS)", "Dockerfiles e manifest K8s com 0 misconfigurações no Trivy"],
    ["CI/CD", "GitHub Actions", "Pipeline de 8 etapas, Actions fixadas por SHA, Dependabot"],
    ["Observabilidade", "Prometheus · Alertmanager · Loki · Alloy · Grafana",
     "Métricas, logs JSON, 11 alertas validados contra ataques simulados, dashboard provisionado"],
], [2.4, 5.4, 9.2], font=8.2)

h2("0.3 Rastreabilidade: requisito do enunciado → onde está atendido")
table(["Subetapa", "Requisito do enunciado", "Onde (seção / arquivo)"], [
    ["1 Pipeline", "Desenho do pipeline CI/CD com foco em segurança", "1.1 (Figura 2) · .github/workflows/devsecops.yml"],
    ["1 Pipeline", "SAST (SonarQube, Semgrep)", "1.2, 1.5, 1.7 · Semgrep + Bandit · .semgrep/"],
    ["1 Pipeline", "SCA (Dependabot, Snyk)", "1.2, 1.5 · pip-audit, npm audit, Dependency Review, .github/dependabot.yml"],
    ["1 Pipeline", "Secret Scanning (GitGuardian, Gitleaks)", "1.2, 1.5, 1.7 · Gitleaks · .gitleaks.toml"],
    ["1 Pipeline", "Container Security (Trivy)", "1.2, 1.5 · Trivy imagem + IaC + SBOM"],
    ["1 Pipeline", "Explicação de como cada etapa reduz riscos", "1.2 (com IDs de ameaça da seção 4.1) e 1.3"],
    ["1 Pipeline", "Como o pipeline seria executado no projeto Ford", "1.4"],
    ["2 Código", "Criptografia local", "2.1 · api/app/crypto.py · mobile/src/services/secureSession.ts"],
    ["2 Código", "Hardening de API (rate limit, validação de entrada, JWT seguro)", "2.2 · api/app/"],
    ["2 Código", "Controle de acesso por perfil (Brigadista, Gestor, Administrador)", "0.1, 2.3 · api/app/security.py"],
    ["2 Código", "Segurança MQTT/TLS para IoT", "2.5 · iot/"],
    ["2 Código", "IaC Security (Dockerfile, Kubernetes YAML)", "2.6 · api/Dockerfile, iot/Dockerfile, infra/k8s/api.yaml"],
    ["2 Código", "Trechos de código, prints, commits, explicações", "Seção 2 inteira; 2.7 (correções reais); 1.5 (saídas das ferramentas)"],
    ["3 Observab.", "Logs estruturados (login, falhas, alterações críticas)", "3.1 · api/app/logging_setup.py · exemplos_logs.jsonl"],
    ["3 Observab.", "Métricas e alertas (API, mobile, IoT, ML)", "3.2 · observability/alert_rules.yml, loki-rules/"],
    ["3 Observab.", "Dashboards (Grafana ou equivalente)", "3.3 (Figura 3) · observability/grafana/"],
    ["3 Observab.", "Plano de resposta: detecção → análise → contenção → erradicação → recuperação", "3.4 (Figura 4 + 6 playbooks)"],
    ["4 Compliance", "Revisão final dos riscos (STRIDE + DevSecOps)", "4.1 e 4.2"],
    ["4 Compliance", "OWASP ASVS · Mobile Top 10 · API Top 10", "4.3, 4.4, 4.5"],
    ["4 Compliance", "LGPD (dados pessoais, telemetria, localização)", "4.6"],
    ["4 Compliance", "Rotinas: dependências, testes, permissões, backup", "4.7"],
    ["4 Compliance", "Documento consolidado + checklist de conformidade", "Este documento + 4.8"],
], [2.2, 7.4, 7.4], font=7.8)

# ================================================================== 1. PIPELINE
h1("1. Pipeline DevSecOps Integrado (peso 3,0)")
p("**Objetivo:** mostrar como a segurança é incorporada ao pipeline, do commit ao deploy. O pipeline está "
  "em `.github/workflows/devsecops.yml` e roda no GitHub Actions a cada push e pull request. Nenhuma etapa "
  "de segurança é apenas informativa: se um *gate* falha, o merge é bloqueado pela proteção de branch e o "
  "deploy não acontece.")
h2("1.1 Desenho do pipeline")
figure(ROOT / "docs" / "img" / "pipeline_devsecops.png",
       "Figura 2 — Pipeline DevSecOps: 8 etapas, com gates bloqueantes do commit até a produção.")

h2("1.2 Etapas e riscos que cada uma reduz")
p("Os IDs entre parênteses remetem às ameaças da revisão STRIDE (seção 4.1) e aos riscos do próprio "
  "ciclo DevSecOps (seção 4.2).", size=9.5)
table(["Etapa", "Ferramenta", "Quando roda", "Bloqueia quando", "Risco reduzido"], [
    ["1 Secret scanning", "Gitleaks 8.30 (regras padrão + 4 do projeto)", "Todo push/PR; todo o histórico",
     "Qualquer segredo encontrado", "Vazamento de JWT_SECRET, chaves AES/HMAC, senha MQTT, chave privada IoT (I-03, S-02)"],
    ["2 SAST", "Semgrep 1.178 (OWASP Top 10, Python, JWT, TS + 7 regras próprias) e Bandit 1.9",
     "Todo push/PR", "Achado ERROR/WARNING; Bandit médio+",
     "JWT sem verificação (S-02), pickle no modelo (T-03), token no AsyncStorage (I-05), MQTT sem TLS (I-06), PII em log (I-04), rota sem auth (E-01)"],
    ["3 SCA", "pip-audit, npm audit, Dependency Review, Dependabot", "Todo PR + semanal",
     "CVE alta/crítica ou licença fora da lista", "Dependência vulnerável ou maliciosa (T-04, OWASP M2)"],
    ["4 Testes", f"pytest ({n_tests} testes: API + IoT)", "Após 1-3", "Qualquer teste falhando",
     "Regressão em autenticação, autorização, validação, cripto, anti-replay e mTLS (S-01..E-01)"],
    ["5 Build + container/IaC", "Trivy 0.74 (2 imagens + config), SBOM CycloneDX, cosign, hash do modelo",
     "Após testes", "CVE alta/crítica corrigível ou misconfig alta/crítica",
     "CVE do SO, container root, K8s inseguro (E-02); imagem trocada (DS-03); modelo trocado (T-03)"],
    ["6 Deploy staging", "azure/login (OIDC), kubectl", "Só na main", "Rollout não fica saudável",
     "Credencial de nuvem fixa no GitHub (DS-02): não existe, o token OIDC dura minutos"],
    ["7 DAST", "OWASP ZAP baseline 0.15", "Após staging", "Regra FAIL (headers, HSTS, CSP)",
     "Configuração insegura só visível em execução (I-04, D-01)"],
    ["8 Produção", "GitHub Environment + cosign verify", "Após DAST + aprovação",
     "Assinatura inválida ou sem aprovação", "Deploy de imagem não verificada ou sem revisão (DS-03, DS-04)"],
], [2.2, 3.5, 2.3, 2.8, 6.2], font=7.6)

h2("1.3 Cobertura por componente (API, mobile, IoT, dados, ML, arquitetura)")
table(["Componente", "Controles do pipeline que o cobrem"], [
    ["API (FastAPI)", "Gitleaks, Semgrep (OWASP/JWT + regras próprias), Bandit, pip-audit, testes de segurança, Trivy imagem, ZAP"],
    ["Mobile (Expo)", "Semgrep TypeScript + regra `asyncstorage-token`, npm audit, Dependabot npm; build do APK (EAS) consome só dependências aprovadas"],
    ["IoT (broker + ingestão)", "Semgrep regra `mqtt-without-tls`, Bandit, pip-audit (iot/requirements.txt), testes de anti-replay e de handshake mTLS, Trivy na imagem da ingestão"],
    ["Dados", "Gitleaks (chaves DATA_KEY/INDEX_KEY), testes de cifra em repouso e rotação de chave, Semgrep `logging-sensitive-field`"],
    ["ML", "Semgrep `pickle-load-model`, hash SHA-256 do modelo calculado no build e exigido pela API em produção, testes de faixas de entrada"],
    ["Arquitetura / infra", "Trivy config (Dockerfiles, Kubernetes), `.trivyignore` com aceite de risco datado, imagens por digest e assinadas"],
], [3.4, 13.6], font=8)

h2("1.4 Como o pipeline é executado no projeto Ford")
for step in [
    "**Desenvolvedor abre um pull request** para `develop`/`main`. A proteção de branch exige os checks "
    "`secret-scan`, `sast`, `sca` e `tests` verdes e 1 revisão de código. PRs vindos de *fork* rodam com o "
    "evento `pull_request`, sem acesso a segredos.",
    "**Etapas 1-3 rodam em paralelo.** Os resultados em SARIF aparecem na aba *Security → Code scanning* do "
    "GitHub, com arquivo e linha. O Dependency Review comenta no PR se alguma dependência nova tiver CVE.",
    "**Etapa 4** instala as dependências fixadas e roda as duas suítes (API e IoT). Os relatórios JUnit ficam "
    "como artefato do workflow.",
    "**Merge na main** dispara a etapa 5: o Trivy analisa Dockerfiles e manifests, as imagens são construídas, "
    "analisadas, recebem SBOM, são publicadas no GHCR e **assinadas** com cosign (keyless). O SHA-256 do "
    "modelo de ML é calculado nesta etapa.",
    "**Etapas 6-7:** deploy em staging pela imagem `@sha256` (imutável), com o hash do modelo injetado em "
    "`MODEL_SHA256`, e varredura DAST com o ZAP.",
    "**Etapa 8:** o Environment `production` pede aprovação de um revisor de segurança; antes do deploy o "
    "pipeline **verifica a assinatura** da imagem. Se ela não foi gerada por este workflow na `main`, o deploy é abortado.",
    "**Agendamento semanal** (segunda, 06:00 UTC) roda tudo de novo, porque CVEs novas aparecem em código parado.",
]:
    bullet(step)

h2("1.5 Execução real das ferramentas (evidência)")
p("Antes de publicar o repositório, rodamos localmente as mesmas ferramentas e versões do pipeline. Os "
  "binários do Gitleaks e do Trivy foram baixados dos releases oficiais e **conferidos pelo SHA-256** do "
  "arquivo de checksums publicado. As saídas completas estão em `docs/evidencias/`.")
table(["Ferramenta", "Escopo", "Resultado", "Evidência"], [
    ["Gitleaks 8.30.1", f"Repositório inteiro (~{gitleaks_kb} KB)",
     f"**{gitleaks_findings} vazamentos.** Controle positivo: 6/6 segredos falsos plantados detectados, incluindo as 4 regras próprias",
     "gitleaks.json, gitleaks-controle.txt"],
    ["Semgrep 1.178.0", f"{semgrep_files} arquivos (API, IoT, mobile, observabilidade)",
     f"**{semgrep_findings} achados** com OWASP Top 10, Python, JWT, TypeScript + regras próprias",
     "semgrep.json"],
    ["Regras próprias Semgrep", "Casos anotados (9 positivos, 6 negativos)",
     "**9/9 detectados, 0 falsos positivos**" if "9/9" in rules_txt else rules_txt.strip(),
     "semgrep-regras.txt"],
    ["Bandit 1.9.4", f"api/app + iot ({bandit_loc} linhas)",
     ("**0 achados**" if bandit_clean else "achados — ver arquivo") +
     f" (após corrigir o B101). {bandit_sim} achados baixos no simulador de tráfego, triados como falso positivo",
     "bandit.txt, bandit-completo.txt"],
    ["pip-audit 2.10.1", "api + iot requirements",
     "**0 vulnerabilidades conhecidas**" if pip_audit_clean else "vulnerabilidades — ver arquivo", "pip-audit.txt"],
    ["Trivy 0.74.0 (config)", "2 Dockerfiles + manifest K8s",
     "**0 misconfigurações** após corrigir o KSV-0013; KSV-0125 com aceite de risco datado", "trivy-config*.txt"],
    ["Trivy 0.74.0 (fs)", "Dependências Python", "**0 CVEs** (segunda opinião do SCA)", "trivy-fs.txt"],
    ["pytest", "API + IoT", f"**{n_tests - n_fail}/{n_tests} aprovados**", "pytest-*.xml / *.txt"],
], [3.0, 3.9, 7.0, 3.1], font=7.8)
p(f"**Triagem do Bandit:** os {bandit_sim} achados de severidade baixa estão em "
  "`observability/simular_trafego.py`, uma ferramenta de teste que não vai para produção. São B311 (`random` "
  "usado para gerar tráfego, não para criptografia) e B105 (senha errada proposital para simular erro de "
  "digitação). Por isso o escopo do Bandit no pipeline é `api/app` e `iot`.", size=9.5)

h2("1.6 Trechos do workflow")
code(snippet(".github/workflows/devsecops.yml", "secret-scan:", "category: gitleaks"),
     "Etapa 1 — Secret scanning em todo o histórico (Action fixada por SHA)")
code(snippet(".github/workflows/devsecops.yml", "- name: Validar as regras próprias", "--sarif --output semgrep.sarif --error"),
     "Etapa 2 — As regras próprias são testadas antes de serem usadas")
code(snippet(".github/workflows/devsecops.yml", "- name: Hash do modelo de ML", "sha256sum api/app/churn_model.json"),
     "Etapa 5 — Hash do modelo de ML, exigido pela API em produção")
code(snippet(".github/workflows/devsecops.yml", "- name: Verificar assinatura antes do deploy", "--certificate-oidc-issuer"),
     "Etapa 8 — Produção só aceita imagem assinada por este workflow na main")

h2("1.7 Regras próprias (ameaças específicas do VIN Share)")
p("Os rulesets públicos não conhecem os riscos do projeto. Por isso escrevemos 7 regras Semgrep e 4 regras "
  "Gitleaks ligadas às ameaças da seção 4.1. Cada regra Semgrep tem casos de teste anotados "
  "(`.semgrep/vinshare-rules.py` e `.ts`), validados no pipeline por `verificar_regras.py`.")
code(snippet(".semgrep/vinshare-rules.yml", "- id: pickle-load-model", "pattern: joblib.load(...)") + "\n\n" +
     snippet(".semgrep/vinshare-rules.yml", "- id: asyncstorage-token", "pattern: AsyncStorage.setItem"),
     "Regras Semgrep do projeto (.semgrep/vinshare-rules.yml)")
code(snippet(".gitleaks.toml", "id = \"vinshare-data-key\"", "tags = [\"crypto\", \"lgpd\"]"),
     "Regra Gitleaks para as chaves de criptografia de dados pessoais (.gitleaks.toml)")

h2("1.8 Proteção do próprio pipeline (supply chain)")
for b in [
    "**Actions fixadas pelo SHA do commit** (a tag fica só no comentário), porque tags podem ser reescritas "
    "(caso `tj-actions/changed-files`, março de 2025). O Dependabot `github-actions` abre PR com o SHA novo.",
    "**Menor privilégio:** `permissions: contents: read` por padrão, cada job pede só o que precisa, e "
    "`persist-credentials: false` no checkout.",
    "**Sem credenciais de nuvem guardadas:** login no Azure por OIDC (`id-token: write`).",
    "**Rastreabilidade:** SBOM CycloneDX, *provenance* `mode=max` e assinatura cosign de cada imagem.",
    "**Exceções com prazo:** o `.trivyignore` só aceita exceção com justificativa, dono e data de expiração. "
    "Quando a data passa, o pipeline volta a falhar.",
]:
    bullet(b)

# ================================================================== 2. CÓDIGO E INFRA
h1("2. Segurança em Código e Infraestrutura (peso 2,5)")
p("**Objetivo:** mostrar práticas de segurança aplicadas diretamente no código e na infraestrutura. Os "
  "trechos abaixo são copiados dos arquivos do repositório. Os controles foram verificados por "
  f"**{n_tests} testes automatizados ({n_tests - n_fail} aprovados)** e pelas ferramentas da seção 1.5. "
  "As correções feitas a partir desses achados estão na seção 2.7.")

h2("2.1 Criptografia local")
p("Os dados pessoais (CPF e telefone) são **cifrados por campo antes de serem gravados**, com AES-256-GCM. "
  "Quem tiver acesso de leitura ao banco ou a um backup vazado vê apenas texto cifrado.")
for b in [
    "**AES-256-GCM** (cifra autenticada): qualquer alteração no texto cifrado é detectada e a leitura falha fechada.",
    "**Nonce aleatório de 96 bits** por operação; **contexto como AAD**, então um valor cifrado não pode ser movido de um campo para outro.",
    "**ID da chave no dado** (`k1:...`): permite rotacionar a chave sem perder os dados antigos (testado).",
    "**Índice cego HMAC-SHA256** para buscar pelo CPF sem guardá-lo em claro, com **separação de domínio** "
    "(`idx:` para CPF, `vin:` para o pseudônimo do VIN) mesmo usando a mesma chave.",
    "**Chaves fora do código** (Azure Key Vault → variável de ambiente); a API não sobe com chave ausente ou de tamanho errado.",
    "**No app mobile:** o token fica no `expo-secure-store` (Android Keystore / iOS Keychain), com "
    "`WHEN_UNLOCKED_THIS_DEVICE_ONLY` e `allowBackup: false`.",
]:
    bullet(b)
code(snippet("api/app/crypto.py", "def encrypt(self, plaintext", "raise DecryptionError"),
     "api/app/crypto.py — cifra e decifra de campos pessoais")

h2("2.2 Hardening da API")
table(["Controle", "Implementação", "Arquivo"], [
    ["JWT seguro", "HS256 fixo (rejeita `alg: none`), segredo de 256 bits ou mais, expiração de 15 min, claims "
     "obrigatórios (iss, aud, exp, nbf, jti), revogação no logout, sem PII no payload", "app/security.py"],
    ["Senhas", "bcrypt com custo 12, mínimo de 12 caracteres, tempo de resposta constante para usuário inexistente", "app/security.py"],
    ["Rate limit", "Login: 5/min por usuário+IP e 20/min por IP. API: 100/min por cliente. Responde 429 com Retry-After", "app/ratelimit.py, routes.py"],
    ["Validação de entrada", "Allow-list Pydantic: tipos, faixas, enums, regex; `extra=forbid` bloqueia mass assignment; "
     "rejeita caracteres de controle (log injection)", "app/schemas.py"],
    ["Tamanho do corpo", "Limite de 64 KB (413), repetido no Ingress", "app/main.py"],
    ["Versão do app", "`X-App-Version` abaixo de `MIN_APP_VERSION` (ou forjada) recebe 426", "app/main.py"],
    ["Cabeçalhos", "HSTS, CSP `default-src 'none'`, nosniff, X-Frame-Options DENY, no-store, Permissions-Policy", "app/main.py"],
    ["Erros", "RFC 9457 (problem+json) sem stack trace; o valor enviado não é devolvido; request_id para o suporte", "app/errors.py"],
    ["Configuração", "Falha no startup com segredo fraco, CORS `*` ou (em produção) sem `MODEL_SHA256`", "app/config.py"],
    ["Exposição", "Swagger desligado em produção; /metrics não é roteado pelo Ingress", "app/main.py, k8s/api.yaml"],
    ["Modelo de ML", "Artefato JSON (sem pickle) com SHA-256 verificado antes de carregar; entradas com faixas", "app/model.py"],
], [2.6, 10.9, 3.5], font=7.9)
code(snippet("api/app/security.py", "def validate(self, token: str)", "if claims[\"jti\"] in self._revoked", extra=3),
     "api/app/security.py — validação do JWT (algoritmo, audience, issuer e claims obrigatórios)")
code(snippet("api/app/routes.py", "for limiter, key, scope in", "raise ApiError(401, \"invalid_credentials\""),
     "api/app/routes.py — login com rate limit duplo, mensagem genérica e auditoria")
code(snippet("api/app/schemas.py", "class ChurnPredictionRequest(Strict):", "connected_vehicle: bool"),
     "api/app/schemas.py — validação por allow-list das entradas do modelo de ML")

h2("2.3 Controle de acesso por perfil")
p("A autorização é feita **sempre no servidor**, em duas camadas: (1) a **permissão do perfil** para a "
  "operação e (2) o **escopo da concessionária** do recurso, que protege contra BOLA/IDOR (OWASP API1). O "
  "perfil que o app recebe no login serve só para montar o menu.")
table(["Permissão", "consultor", "gestor", "admin", "Endpoints"], [
    ["leads:read", "✔ (própria)", "✔ (própria)", "✔ (todas)", "GET /dealers/{id}/leads, GET /leads/{id}"],
    ["leads:update", "✔ (própria)", "✔ (própria)", "✔ (todas)", "PATCH /leads/{id}"],
    ["service_share:read", "—", "✔ (própria)", "✔ (todas)", "GET /dealers/{id}/service-share"],
    ["predictions:create", "—", "✔", "✔", "POST /predictions/churn"],
    ["audit:read", "—", "—", "✔", "GET /admin/audit-events"],
    ["users:manage", "—", "—", "✔", "(gestão de usuários)"],
], [3.2, 2.1, 2.1, 1.9, 7.7], font=8.5)
code(snippet("api/app/security.py", "PERMISSIONS: dict[Role, set[str]] = {", "\"dealers:all\"}", extra=1),
     "api/app/security.py — matriz de permissões (menor privilégio)")
code(snippet("api/app/routes.py", "def _get_lead_scoped(", "return lead"),
     "api/app/routes.py — proteção contra BOLA: lead de outra concessionária responde 404, sem revelar que existe")

h2("2.4 Mobile: sessão segura e versão mínima")
p("O app guarda o token só no armazenamento seguro do sistema, usa apenas HTTPS e envia a própria versão. "
  "Quando uma versão tem falha de segurança conhecida, basta subir `MIN_APP_VERSION` na API: os aparelhos "
  "desatualizados recebem 426 e o app pede a atualização (OWASP Mobile M8).")
code(snippet("mobile/src/services/secureSession.ts", "// 401 só significa", "throw new UpdateRequiredError", extra=1),
     "mobile/src/services/secureSession.ts — 401 com token = sessão expirada; 426 = atualização obrigatória")
code(snippet("api/app/main.py", "elif app_version is not None and outdated:", "\"Versão do aplicativo não suportada"),
     "api/app/main.py — a API recusa versões antigas ou forjadas do app")

h2("2.5 Segurança MQTT/TLS para IoT")
for b in [
    "**Somente TLS (porta 8883)**, versão mínima 1.2 e cifras ECDHE+AES-GCM. A porta 1883 não existe.",
    "**mTLS:** cada módulo telemático tem certificado X.509 próprio; o CN é o **pseudônimo do VIN** e vira o usuário na ACL.",
    "**ACL por dispositivo:** o veículo só publica em `vehicles/<seu-id>/telemetry`; a ingestão só lê.",
    "**Revogação:** `gerar_certificados.ps1 -Revogar <id>` revoga o certificado e regenera a CRL (playbook PB-04).",
    "**Ingestão desconfia da telemetria** (OWASP API10): tópico, 4 KB, esquema estrito, janela de 5 min, "
    "**anti-replay** por `msg_id` e odômetro que não pode regredir. Cada rejeição vira métrica e log.",
    "**Minimização:** só odômetro, códigos de falha e vida do óleo. Não há GPS.",
]:
    bullet(b)
code(snippet("iot/mosquitto/mosquitto.conf", "per_listener_settings false", "acl_file"),
     "iot/mosquitto/mosquitto.conf — broker com TLS obrigatório, mTLS e CRL")
code(snippet("iot/mosquitto/acl.conf", "pattern write", "topic read vehicles/+/telemetry"),
     "iot/mosquitto/acl.conf — isolamento entre veículos")
code(snippet("iot/ingestion.py", "if data[\"ts\"] < now - self.max_age:", "return self._reject(\"odometro_regrediu\", device)"),
     "iot/ingestion.py — janela de tempo, anti-replay e coerência do odômetro")
p("**Handshake mTLS testado de verdade** (`iot/tests/test_mtls.py`): o teste gera uma PKI, sobe um servidor "
  "TLS com a política do broker e conecta com o mesmo contexto TLS do dispositivo. Casos: dispositivo "
  "válido conecta (o CN chega ao broker); certificado de **CA falsa com o mesmo CN** é recusado; cliente "
  "sem certificado é recusado; broker sem SAN correspondente ou com hostname errado é recusado pelo "
  "dispositivo (proteção contra MITM).", size=9.5)

h2("2.6 IaC Security (Dockerfile e Kubernetes)")
table(["Prática", "Onde", "Ameaça mitigada"], [
    ["Build multi-stage; o runtime não tem pip nem compilador", "api/Dockerfile", "Menor superfície de ataque"],
    ["Usuário não-root (UID 10001) e código somente leitura", "api/ e iot/Dockerfile", "Escalada de privilégio"],
    ["Nenhum segredo na imagem; certificados e chaves montados em runtime", "Dockerfiles, compose, k8s", "Vazamento pela imagem"],
    ["runAsNonRoot, readOnlyRootFilesystem, drop ALL, seccomp RuntimeDefault", "k8s/api.yaml", "Escape de container"],
    ["requests/limits de CPU e memória", "k8s/api.yaml, compose", "Negação de serviço"],
    ["NetworkPolicy default-deny (entrada só do Ingress/Prometheus; saída só DB e DNS)", "k8s/api.yaml", "Movimento lateral"],
    ["automountServiceAccountToken: false", "k8s/api.yaml", "Abuso da API do Kubernetes"],
    ["Imagem sempre por digest @sha256 e assinada", "pipeline + k8s", "Troca de imagem"],
    ["Ingress: TLS 1.2/1.3, 64 KB, 20 rps, /metrics bloqueado", "k8s/api.yaml", "DoS e exposição de métricas"],
    ["Compose: portas só em 127.0.0.1, read_only, no-new-privileges, cap_drop ALL", "infra/docker-compose.yml", "Exposição local"],
], [8.2, 3.6, 5.2], font=7.9)
p("**Resultado do Trivy 0.74 (config):** 2 Dockerfiles com 0 achados. No manifest K8s, 99 de 100 "
  "checagens passaram na primeira execução. O achado KSV-0013 foi corrigido (seção 2.7) e o KSV-0125 "
  "(registry fora da lista padrão) ficou como aceite de risco com prazo no `.trivyignore`.", size=9.5)
code(snippet("infra/k8s/api.yaml", "allowPrivilegeEscalation: false", "limits:   { cpu: 500m, memory: 256Mi }"),
     "infra/k8s/api.yaml — contexto de segurança do container")

h2("2.7 Evidências: testes e correções reais")
labels = {
    "TestAuthentication": ("API · Autenticação / JWT", "login, expiração, assinatura, alg=none, audience/issuer, adulteração, logout"),
    "TestAuthorization": ("API · Autorização (RBAC + BOLA)", "perfis, outra concessionária, 404 em lead alheio, mass assignment"),
    "TestHardening": ("API · Validação e hardening", "enum/tamanho/log injection, faixas ML, 429, 413, 426, headers, CORS, erro 500, hash obrigatório"),
    "TestDataProtection": ("API · Criptografia e dados pessoais", "PII cifrada, CPF mascarado, adulteração, troca de campo, rotação, integridade do modelo"),
    "TestAuditLogging": ("API · Logs de auditoria", "eventos de login sem segredo no log, 403 auditado, alteração crítica"),
    "test_ingestion": ("IoT · Ingestão", "anti-replay, janela de tempo, esquema, tópico, 4 KB, odômetro, métricas"),
    "test_mtls": ("IoT · Handshake mTLS", "dispositivo válido, CA falsa, sem certificado, sem SAN, hostname errado"),
}
table(["Categoria", "Testes", "O que é verificado", "Resultado"],
      [[labels[c][0], str(by_class[c]), labels[c][1],
        "PASSOU" if all(r == "PASSOU" for cc, _, r in all_cases if cc == c) else "FALHOU"] for c in labels],
      [4.0, 1.3, 9.5, 2.2], font=7.8)

h3("Correções reais encontradas pelos testes, pelas ferramentas e pela revisão")
p("Tudo abaixo foi encontrado e corrigido durante esta Sprint. É a evidência de que o ciclo DevSecOps "
  "funciona: a ferramenta ou o teste acha, a correção entra e a verificação passa a garantir.", size=9.5)
table(["#", "Encontrado por", "Problema", "Correção"], [
    ["1", "Teste (pytest)", "Campo `note` aceitava `\\r\\n` (log injection): a regex de controle deixava 0x0A e 0x0D de fora",
     "Faixa 0x0A-0x1F bloqueada (schemas.py); 43/44 → 44/44"],
    ["2", "Semgrep (execução local)", "Arquivo de regras era YAML inválido (`:` sem aspas): no CI a etapa 2 quebraria",
     "Padrões em bloco literal + casos de teste + `verificar_regras.py` no pipeline"],
    ["3", "Semgrep (casos de teste)", "Regra `route-without-auth-dependency` não reconhecia parâmetro tipado (`p: T = Depends()`)",
     "Padrão adicional; 9/9 casos detectados"],
    ["4", "Bandit B101", "`assert` validando MIN_APP_VERSION (some com `python -O`)", "Checagem explícita com ConfigError"],
    ["5", "Trivy KSV-0013", "Placeholder de imagem sem tag/digest no manifest K8s", "Placeholder `@sha256:` válido; `sed` do pipeline ajustado"],
    ["6", "Trivy KSV-0125", "Registry (GHCR) fora da lista padrão", "Aceite de risco com dono e expiração (31/03/2027)"],
    ["7", "Revisão de versões", "Actions e ferramentas defasadas e fixadas por tag (ex.: artefatos do Trivy 0.57 nem existem mais)",
     "Versões atuais, fixadas por SHA de commit"],
    ["8", "Simulação de ataques", f"Alertas não disparariam: força bruta chegou a só {old_bruteforce_peak:.0f} falhas "
     f"(o rate limit bloqueia antes) contra limiar 30; BOLA chegou a {old_bola_peak:.0f} contra 20",
     "Regra soma falhas + bloqueios; limiar de 403 = 5; todos os ataques validados (3.3)"],
    ["9", "Teste mTLS", "Certificados sem AKI/SKI são recusados pela verificação X.509 estrita (Python 3.13+)",
     "SKI/AKI explícitos no gerar_certificados.ps1"],
    ["10", "Revisão", "Broker sem SAN (verificação de hostname falharia) e sem CRL gerada (broker não subiria)",
     "SAN localhost/mosquitto, CRL e comando de revogação no script"],
    ["11", "Revisão", "App mostrava \"sessão expirada\" para senha errada (401 no login)", "401 só limpa sessão se havia token"],
    ["12", "Revisão", "Alertas de IoT/mobile usavam métricas que não existem", "Serviço de ingestão com métricas reais + métrica de versão do app"],
    ["13", "Revisão", "Hash do modelo de ML era opcional mesmo em produção", "ConfigError em prod + hash gerado no pipeline"],
    ["14", "Revisão", "Log da ingestão IoT era JSON aninhado (filtro do Loki não acharia o evento)", "Log plano no mesmo formato da API"],
], [0.7, 2.8, 7.2, 6.3], font=7.5)
callout("**Para anexar antes da entrega:** (1) publicar o repositório, fazer um commit por correção da tabela "
        "acima e colar aqui os links dos commits; (2) colar o print da aba *Actions* com o workflow "
        "`devsecops` verde.")

# ================================================================== 3. OBSERVABILIDADE
h1("3. Observabilidade, Monitoramento e Resposta (peso 2,0)")
p("**Objetivo:** mostrar como o sistema detecta, registra e responde a incidentes. As métricas e os logs "
  "abaixo são **reais**. Foram gerados executando a API e a ingestão IoT com `observability/simular_trafego.py`, "
  "que simula 2 horas de uso normal com cinco ataques injetados: força bruta, sondagem BOLA, tokens forjados, "
  "app desatualizado e replay de telemetria.")

h2("3.1 Plano de monitoramento: logs estruturados")
p("API e ingestão escrevem um evento JSON por linha no stdout. O Grafana Alloy envia os eventos ao **Loki**. "
  "Todo evento tem timestamp UTC, nível, nome do evento e, na API, o **request_id**, que liga o log à "
  "resposta de erro recebida pelo usuário. O `RedactingFilter` troca por `[REDACTED]` campos chamados "
  "senha, token, CPF, telefone, e-mail ou VIN.")
table(["Evento", "Quando é gerado", "Campos principais", "Na simulação"], [
    ["auth.login.success", "Login aceito", "user_id, role, dealer_id, client_ip", str(event_counts["auth.login.success"])],
    ["auth.login.failure", "Senha errada / usuário inexistente", "username, client_ip", str(event_counts["auth.login.failure"])],
    ["auth.login.blocked", "Rate limit de login disparado", "username, scope, client_ip", str(event_counts["auth.login.blocked"])],
    ["auth.token_rejected", "JWT inválido, expirado ou revogado", "reason, client_ip, path", str(event_counts["auth.token_rejected"])],
    ["authz.denied", "Perfil sem permissão ou fora do escopo", "user_id, role, permission, target_dealer", str(event_counts["authz.denied"])],
    ["lead.status_changed", "Alteração crítica de dado de negócio", "user_id, lead_id, old_status, new_status", str(event_counts["lead.status_changed"])],
    ["iot.telemetry.rejected", "Telemetria recusada pela ingestão", "reason, device_id", str(event_counts["iot.telemetry.rejected"])],
    ["http.access", "Toda requisição (inclui 426 do app antigo)", "method, route, status, duration_ms, user_id", str(event_counts["http.access"])],
    ["app.error", "Exceção não tratada (500)", "exc_type, stack (só no log)", str(event_counts.get("app.error", 0))],
], [3.5, 4.4, 6.4, 2.7], font=7.9)
p("Retenção: eventos de auditoria ficam **1 ano** em armazenamento somente-anexo. Logs de acesso ficam "
  "**6 meses**, o mínimo exigido pelo Marco Civil da Internet (art. 15).", size=9.5)

h3("Exemplos reais de log (extraídos de observability/exemplos_logs.jsonl)")
examples = [
    first_log("auth.login.failure", client_ip="203.0.113.50"),
    first_log("auth.login.blocked"),
    first_log("authz.denied", user_id="u-200"),
    first_log("auth.token_rejected"),
    first_log("lead.status_changed"),
    first_log("iot.telemetry.rejected", reason="replay"),
    first_log("http.access", status=426),
]
code("\n".join(json.dumps(e, ensure_ascii=False) for e in examples), size=6.5)

h2("3.2 Métricas e alertas (API, mobile, IoT e ML)")
p("A API e a ingestão expõem métricas Prometheus, acessíveis só pela rede interna. As regras ficam em "
  "`observability/alert_rules.yml` (métricas) e `observability/loki-rules/iot.yml` (logs do broker). O "
  "Alertmanager envia *warning* ao canal do time no Teams e *critical* ao plantão de segurança. As URLs "
  "de webhook ficam em arquivos de segredo.")
table(["Área", "Fonte / métrica", "Alerta", "Condição", "Sev.", "Playbook"], [
    ["API", "auth_login_total{failure} + rate_limited_total{login}", "ForcaBrutaLogin", "> 30 em 5 min", "crítico", "PB-01"],
    ["API", "vinshare_authz_denied_total", "SondagemDeAutorizacao", "> 5 em 5 min", "crítico", "PB-02"],
    ["API", "vinshare_auth_token_rejected_total", "TokensForjadosOuExpirados", "> 20 inválidos em 5 min", "crítico", "PB-03"],
    ["API", "vinshare_http_requests_total", "ErrosServidor", "5xx > 2% por 5 min", "crítico", "—"],
    ["API", "vinshare_http_request_duration_seconds", "LatenciaAlta", "p95 > 500 ms por 10 min", "aviso", "—"],
    ["Mobile", "vinshare_mobile_requests_total{bloqueada}", "AppDesatualizadoEmUso", "> 10 (426) em 15 min", "aviso", "—"],
    ["IoT", "vinshare_iot_messages_total{≠aceita}", "TelemetriaRejeitada", "> 10 em 10 min", "crítico", "PB-04"],
    ["IoT", "vinshare_iot_last_accepted_timestamp", "TelemetriaParou", "nada aceito há 15 min", "aviso", "—"],
    ["IoT", "Log do broker (LogQL)", "FalhaAutenticacaoMQTT", "> 20 recusas TLS em 10 min", "crítico", "PB-04"],
    ["ML", "vinshare_ml_churn_score (histograma)", "DriftScoreChurn", "> 50% com risco alto em 1 h", "aviso", "PB-05"],
    ["ML", "vinshare_ml_predictions_total", "VolumePredicoesAnomalo", "> 500 em 5 min", "aviso", "PB-05"],
], [1.3, 5.0, 3.7, 3.2, 1.4, 2.4], font=7.4)
p("**Mobile (crashes):** a estabilidade do app é monitorada no Firebase Crashlytics, com alerta nativo no "
  "console quando os usuários sem crash ficam abaixo de 99%. Essa métrica não passa pelo Prometheus.", size=9.5)

h2("3.3 Dashboard e validação dos alertas")
p("O dashboard `observability/grafana/vinshare-security.json` é provisionado automaticamente no Grafana "
  "pelo `docker-compose`, com Prometheus e Loki como fontes. A Figura 3 mostra os mesmos painéis, montados "
  "com as métricas coletadas na simulação. O último painel mostra quando cada regra de alerta teria disparado.")
figure(ROOT / "docs" / "img" / "painel_monitoramento.png",
       "Figura 3 — Painel de segurança e operação com as métricas reais da simulação (2 h, intervalos de 5 min).")
p("**Validação dos limiares** (`observability/avaliar_alertas.py`): cada regra foi avaliada sobre a série "
  "da simulação. Critério: todo ataque dispara o alerta do seu playbook e o tráfego normal não dispara nada.", size=9.5)
table(["Alerta", "Limiar", "Pico observado", "Ataque (min)", "Disparou em (min)", "Resultado"],
      [[r["alerta"], r["limiar"], f"{r['pico']:.3g}", ", ".join(str(t * 5) for t in r["ataque"]) or "—",
        ", ".join(str(t * 5) for t in r["disparos"]) or "—",
        ("detectou" if r["detectou"] else "sem disparo (correto)" if not r["ataque"] else "NÃO detectou")
        + ("" if not r["falsos_positivos"] else " + falso positivo")]
       for r in alert_report],
      [3.7, 4.5, 1.7, 2.0, 2.4, 2.7], font=7.4)
callout("**Para anexar:** rodar `docker compose up` em `infra/` (ver README), abrir o Grafana em "
        "http://localhost:3000 (pasta *VIN Share*) e colar aqui o print do dashboard.")

h2("3.4 Plano de resposta a incidentes")
figure(ROOT / "docs" / "img" / "resposta_incidentes.png",
       "Figura 4 — Fluxo de resposta a incidentes, baseado no NIST SP 800-61.")
h3("Classificação de severidade")
table(["Sev.", "Exemplo", "1ª resposta", "Contenção", "Quem é acionado"], [
    ["P1", "Vazamento de dados pessoais; chave comprometida", "15 min", "até 4 h", "Plantão, líder de segurança, DPO, jurídico"],
    ["P2", "Ataque ativo (força bruta, BOLA, replay IoT) já contido pelos controles", "1 h", "até 24 h", "Plantão e líder técnico"],
    ["P3", "Vulnerabilidade alta sem exploração; alerta de drift", "1 dia útil", "até 7 dias", "Time do produto"],
    ["P4", "Achado baixo, melhoria", "5 dias úteis", "próxima sprint", "Time do produto"],
], [1.2, 5.6, 2.1, 2.3, 5.8], font=8)

h3("Playbooks")
table(["ID", "Cenário", "Detecção → Análise", "Contenção", "Erradicação → Recuperação"], [
    ["PB-01", "Força bruta / credential stuffing", "ForcaBrutaLogin → IPs e usuários em auth.login.failure/blocked",
     "Rate limit já bloqueia; bloquear IP/ASN no Ingress/WAF; forçar troca de senha das contas-alvo",
     "Avaliar MFA para gestores; monitorar 72 h; desbloquear contas verificadas"],
    ["PB-02", "Sondagem BOLA / abuso interno", "SondagemDeAutorizacao → user_id e target_dealer em authz.denied",
     "Desativar usuário e revogar tokens (jti)", "Revisar permissões da concessionária; conferir se houve 200 indevido; comunicar gestor/RH"],
    ["PB-03", "Token forjado / JWT_SECRET vazado", "TokensForjadosOuExpirados → padrão de tokens e IPs",
     "Rotacionar JWT_SECRET no Key Vault (invalida todos os tokens)", "Achar a origem do vazamento (Gitleaks, logs); redeploy; novo login para todos"],
    ["PB-04", "Dispositivo IoT comprometido / replay", "TelemetriaRejeitada ou FalhaAutenticacaoMQTT → device_id e reason",
     "`gerar_certificados.ps1 -Revogar <id>` (CRL); descartar telemetria do período", "Reemitir certificado; investigar o firmware; reprocessar dados válidos"],
    ["PB-05", "Abuso ou drift do modelo de ML", "DriftScoreChurn / VolumePredicoesAnomalo",
     "Rate limit por usuário nas predições; congelar a versão do modelo", "Investigar envenenamento; retreinar com dados validados; publicar novo hash"],
    ["PB-06", "Vazamento de dados pessoais", "Qualquer fonte (alerta, relato, terceiro)",
     "Isolar o componente; preservar evidências (logs, snapshots)",
     "Avaliar risco aos titulares com o DPO; comunicar ANPD e titulares (LGPD art. 48)"],
], [1.2, 2.9, 4.0, 4.3, 4.6], font=7.4)
callout("**LGPD:** incidente com dado pessoal que possa gerar risco ou dano relevante aos titulares deve ser "
        "comunicado à ANPD e aos titulares (art. 48). O Regulamento de Comunicação de Incidente de Segurança "
        "(Resolução CD/ANPD nº 15/2024) fixa o prazo de **3 dias úteis**. O playbook PB-06 aciona o DPO já na fase de análise.")

# ================================================================== 4. COMPLIANCE
h1("4. Compliance, Riscos e Segurança Contínua (peso 2,5)")
h2("4.1 Revisão final dos riscos (STRIDE)")
p("Ameaças por componente, com o controle implementado, a verificação contínua no pipeline e o risco residual.")
stride = [
    ["S-01", "API", "S", "Força bruta / credential stuffing no login", "bcrypt 12, rate limit duplo, mensagem genérica, alerta", "Etapa 4 (testes), alerta validado", "Baixo"],
    ["S-02", "API", "S", "JWT forjado (alg=none, chave errada, claims)", "HS256 fixo, iss/aud/exp/jti obrigatórios", "Etapa 4, Semgrep (2 regras)", "Baixo"],
    ["S-03", "IoT", "S", "Dispositivo falso publicando telemetria", "mTLS, certificado por veículo, CRL, ACL", "Etapa 4 (teste mTLS)", "Baixo"],
    ["T-01", "IoT", "T", "Telemetria adulterada ou repetida (replay)", "TLS; ingestão com janela de tempo, msg_id e odômetro", "Etapa 4 (23 testes), alerta", "Baixo"],
    ["T-02", "API", "T", "Mass assignment (enviar dealer_id/role)", "extra=forbid nos schemas", "Etapa 4 (testes)", "Baixo"],
    ["T-03", "ML", "T", "Modelo trocado ou envenenado", "JSON + SHA-256 obrigatório em prod; alerta de drift", "Etapas 4 e 5, Semgrep (pickle)", "Médio"],
    ["T-04", "CI/CD", "T", "Dependência, imagem ou Action maliciosa", "SCA, SBOM, cosign, pin por SHA", "Etapas 3 e 5", "Médio"],
    ["R-01", "API", "R", "Usuário nega ter alterado um lead", "Log de auditoria com user_id e request_id (1 ano)", "Etapa 4 (testes)", "Baixo"],
    ["I-01", "Dados", "I", "Vazamento do banco ou de backup", "AES-256-GCM por campo, TDE, backup cifrado", "Etapa 4 (testes)", "Baixo"],
    ["I-02", "API", "I", "BOLA: ler leads de outra concessionária", "Escopo por concessionária, 404", "Etapa 4, alerta validado", "Baixo"],
    ["I-03", "Repo", "I", "Segredo commitado", "Key Vault + env; Gitleaks em todo o histórico", "Etapa 1", "Baixo"],
    ["I-04", "API", "I", "Stack trace ou dado sensível no erro/log", "problem+json, RedactingFilter", "Etapa 4, Semgrep", "Baixo"],
    ["I-05", "Mobile", "I", "Token extraído do aparelho", "SecureStore, allowBackup false, HTTPS obrigatório", "Semgrep (AsyncStorage)", "Médio"],
    ["I-06", "IoT", "I", "Telemetria interceptada na rede", "TLS obrigatório, porta 1883 inexistente", "Semgrep (MQTT sem TLS)", "Baixo"],
    ["D-01", "API", "D", "Flood de requisições / payload grande", "Rate limit, 64 KB, limit-rps, limites K8s", "Etapa 4, DAST", "Médio"],
    ["D-02", "IoT", "D", "Flood MQTT", "Pacote de 4 KB, limites de fila, ACL", "Revisão de config", "Médio"],
    ["E-01", "API", "E", "Consultor acessando função de admin", "RBAC no servidor por permissão", "Etapa 4, Semgrep (rota sem auth)", "Baixo"],
    ["E-02", "Infra", "E", "Escape de container / escalada no cluster", "Não-root, drop ALL, read-only, seccomp, sem SA token", "Etapa 5 (Trivy config)", "Baixo"],
]
table(["ID", "Comp.", "STRIDE", "Ameaça", "Controle implementado", "Verificação contínua", "Residual"],
      stride, [1.0, 1.3, 1.2, 3.9, 4.8, 3.0, 1.8], font=7.3)
p("**Riscos residuais médios** (aceitos para esta fase, com plano): T-03 inclui envenenamento dos dados de "
  "treino; T-04 cobre supply chain, que nunca chega a zero; I-05 considera aparelhos com root/jailbreak "
  "(plano: detecção de root e certificate pinning no build nativo); D-01 e D-02 dependem de proteção DDoS "
  "de rede (Azure DDoS Protection / WAF).", size=9.5)

h2("4.2 Riscos do ciclo DevSecOps")
table(["ID", "Risco", "Controle", "Residual"], [
    ["DS-01", "PR de fork rouba segredos do pipeline", "Evento `pull_request` (sem segredos); nenhum `pull_request_target`", "Baixo"],
    ["DS-02", "Credencial de nuvem vazada do CI", "OIDC com token de curta duração; nenhum segredo de nuvem guardado", "Baixo"],
    ["DS-03", "Imagem adulterada entre build e deploy", "Deploy por digest + cosign verify com identidade do workflow", "Baixo"],
    ["DS-04", "Bypass dos gates (merge direto / admin)", "Branch protection para todos (inclui admins) + Environment com revisor", "Baixo"],
    ["DS-05", "Action de terceiro comprometida", "Pin por SHA, permissões mínimas, Dependabot", "Médio"],
    ["DS-06", "Fadiga de alertas / falso negativo das ferramentas", "Regras próprias testadas, limiares validados por simulação, triagem documentada", "Médio"],
    ["DS-07", "Exceção de segurança esquecida", "`.trivyignore` com dono e data de expiração", "Baixo"],
], [1.3, 5.0, 8.7, 2.0], font=7.8)

h2("4.3 Mapeamento OWASP ASVS 4.0.3")
table(["Req.", "Descrição resumida", "Como atendemos", "Status"], [
    ["2.1.1", "Senhas com 12+ caracteres", "LoginRequest.min_length=12", "Implementado e testado"],
    ["2.2.1", "Controles anti-automação", "Rate limit por usuário+IP e por IP", "Implementado e testado"],
    ["2.4.1", "Hash de senha resistente (bcrypt/argon2...)", "bcrypt custo 12", "Implementado"],
    ["3.3.1", "Logout invalida a sessão", "Revogação do jti", "Implementado e testado"],
    ["3.5.3", "Tokens stateless assinados e protegidos contra adulteração", "HS256 fixo, claims obrigatórios, sem alg=none", "Implementado e testado"],
    ["4.1.1 / 4.1.3", "Controle de acesso no servidor, com menor privilégio", "RBAC em dependências FastAPI", "Implementado e testado"],
    ["4.2.1", "Proteção contra IDOR", "Escopo por concessionária", "Implementado e testado"],
    ["5.1.2 / 5.1.3", "Proteção contra mass assignment; validação por allow-list", "Pydantic extra=forbid, enums, faixas", "Implementado e testado"],
    ["6.2.1 / 6.2.2", "Criptografia falha de forma segura, com algoritmos aprovados", "AES-256-GCM, DecryptionError", "Implementado e testado"],
    ["6.4.1", "Gestão de chaves", "Azure Key Vault, ID de chave para rotação", "Implementado (código) / Key Vault planejado"],
    ["7.1.1", "Sem credenciais ou tokens em log", "RedactingFilter + teste", "Implementado e testado"],
    ["7.1.3", "Registrar eventos de segurança", "auth.*, authz.denied, lead.status_changed, iot.*", "Implementado e testado"],
    ["7.4.1", "Erro genérico com ID para o suporte", "problem+json com request_id", "Implementado e testado"],
    ["8.3.1", "Dado sensível no corpo/cabeçalho, nunca na URL", "Login por POST no corpo; token no header", "Implementado"],
    ["9.1.1 / 9.1.2", "TLS em todas as conexões, cifras fortes", "Ingress TLS 1.2/1.3, MQTT 8883 (mTLS testado)", "Implementado e testado (IoT)"],
    ["14.2.1", "Componentes atualizados, sem CVE conhecida", "Dependabot + pip-audit + Trivy (0 CVEs hoje)", "Executado"],
    ["14.3.2", "Modos de debug desligados em produção", "Swagger/OpenAPI off em APP_ENV=prod", "Implementado"],
    ["14.4.3-14.4.7", "Cabeçalhos de segurança", "CSP, nosniff, HSTS, frame-ancestors", "Implementado e testado"],
], [2.2, 5.0, 5.6, 4.2], font=7.7)

h2("4.4 OWASP API Security Top 10 (2023)")
table(["Risco", "Controle no VIN Share", "Status"], [
    ["API1 Broken Object Level Authorization", "Escopo de concessionária em todo recurso; 404 para recurso alheio", "Implementado e testado"],
    ["API2 Broken Authentication", "JWT de 15 min com revogação; bcrypt; rate limit no login", "Implementado e testado"],
    ["API3 Broken Object Property Level Authorization", "extra=forbid na entrada; CPF mascarado na saída", "Implementado e testado"],
    ["API4 Unrestricted Resource Consumption", "Rate limit, corpo de 64 KB, limites K8s e Ingress, limit ≤ 500", "Implementado e testado"],
    ["API5 Broken Function Level Authorization", "Permissão por função (audit:read só admin)", "Implementado e testado"],
    ["API6 Unrestricted Access to Sensitive Business Flows", "Alerta de volume de predições (extração do modelo); rate limit", "Implementado"],
    ["API7 Server Side Request Forgery", "A API não busca URLs vindas do usuário; egress restrito por NetworkPolicy", "Não aplicável por design"],
    ["API8 Security Misconfiguration", "Headers, CORS allow-list, config que falha no startup, Trivy config", "Implementado e testado"],
    ["API9 Improper Inventory Management", "Versão /api/v1; docs off em prod; staging separado", "Implementado"],
    ["API10 Unsafe Consumption of APIs", "Telemetria tratada como não confiável (esquema, tempo, anti-replay)", "Implementado e testado"],
], [5.4, 8.2, 3.4], font=7.8)

h2("4.5 OWASP Mobile Top 10 (2024)")
table(["Risco", "Controle no app VIN Share", "Status"], [
    ["M1 Improper Credential Usage", "Nenhuma credencial no app; só o token do usuário no SecureStore", "Implementado (código)"],
    ["M2 Inadequate Supply Chain Security", "npm audit + Dependabot npm + Dependency Review", "Configurado no pipeline"],
    ["M3 Insecure Authentication/Authorization", "Autorização só na API; 401 com token limpa a sessão", "Implementado (código)"],
    ["M4 Insufficient Input/Output Validation", "Validação na API (allow-list); o app não monta HTML com dados da API", "Implementado"],
    ["M5 Insecure Communication", "Somente HTTPS; usesCleartextTraffic=false; pinning no build nativo", "Implementado / pinning planejado"],
    ["M6 Inadequate Privacy Controls", "Sem localização/contatos; CPF mascarado; nada de PII em log", "Implementado (código)"],
    ["M7 Insufficient Binary Protections", "ProGuard/R8 e shrink de recursos no release", "Configurado"],
    ["M8 Security Misconfiguration", "allowBackup=false; permissões bloqueadas; versão mínima forçada (426)", "Implementado e testado (API)"],
    ["M9 Insecure Data Storage", "expo-secure-store; regra Semgrep contra AsyncStorage (testada)", "Implementado"],
    ["M10 Insufficient Cryptography", "Criptografia no servidor com AES-GCM; o app não implementa cripto própria", "Implementado"],
], [5.4, 8.2, 3.4], font=7.8)

h2("4.6 LGPD: dados pessoais, telemetria e localização")
table(["Artigo", "Exigência", "Aplicação no VIN Share", "Status"], [
    ["Art. 6º, I-III", "Finalidade, adequação e necessidade", "Telemetria limitada a odômetro, códigos de falha e vida do óleo (esquema estrito rejeita campos extras); CPF sempre mascarado", "Implementado e testado"],
    ["Art. 7º", "Base legal", "Execução de contrato/garantia para manutenção; legítimo interesse (com LIA e opt-out) para lembretes; **consentimento** para localização e ofertas", "Planejado (texto jurídico)"],
    ["Art. 9º", "Transparência", "Aviso de privacidade no app, explicando telemetria e modelo de churn", "Planejado"],
    ["Art. 13, §4º", "Pseudonimização", "VIN pseudonimizado (HMAC) no IoT, analytics e ML; continua tratado como dado pessoal", "Implementado"],
    ["Art. 16", "Eliminação ao fim do tratamento", "Telemetria bruta: 24 meses; auditoria: 1 ano; acesso: 6 meses", "Planejado (job de expurgo)"],
    ["Art. 18", "Direitos do titular", "Fluxos de acesso, correção, eliminação e revogação de consentimento pelo app/SAC", "Planejado"],
    ["Art. 37 / 38", "Registro das operações e RIPD", "Inventário de dados + RIPD, pois telemetria veicular é tratamento de risco elevado", "Planejado"],
    ["Art. 46", "Medidas de segurança", "AES-256-GCM, TLS/mTLS, RBAC, logs, pipeline DevSecOps", "Implementado"],
    ["Art. 48", "Comunicação de incidente", "Playbook PB-06; ANPD em 3 dias úteis (Res. CD/ANPD nº 15/2024)", "Implementado (processo)"],
], [2.2, 3.4, 8.2, 3.2], font=7.7)
p("**Localização:** não é dado sensível pela LGPD (art. 5º, II), mas revela hábitos e endereço do titular. "
  "Por isso a solução **não coleta GPS**, e o esquema da ingestão rejeita qualquer campo extra. Se uma "
  "funcionalidade futura precisar de localização (ex.: concessionária mais próxima), a coleta será pontual, "
  "com consentimento e precisão reduzida (cidade/UF).", size=9.5)

h2("4.7 Plano de segurança contínua")
table(["Rotina", "Frequência", "Como / ferramenta", "Responsável", "Evidência"], [
    ["Revisão de dependências", "A cada PR + semanal",
     "Dependabot (pip, npm, Docker, compose, Actions), pip-audit, npm audit, Trivy. SLA: crítica 7 dias, alta 30, média 90",
     "Tech lead", "PRs do Dependabot; aba Security"],
    ["Testes de segurança", "A cada PR", "Gitleaks, Semgrep (+ teste das regras), Bandit, pytest API + IoT", "Todo o time", "Checks do PR"],
    ["", "Semanal / a cada deploy", "Re-scan agendado; ZAP baseline em staging; revisão das exceções do `.trivyignore`", "DevOps", "Artefatos do workflow"],
    ["", "A cada mudança de alerta", "Reexecutar simulação + `avaliar_alertas.py` (todo ataque dispara, nada de falso positivo)", "Segurança", "Saída do avaliador"],
    ["", "Semestral", "Pentest (API, app e broker) + revisão STRIDE de novas features", "Segurança", "Relatório de pentest"],
    ["Auditoria de permissões", "Mensal", "Usuários ativos x concessionárias/RH; desligados removidos em 24 h; revisão de authz.denied",
     "Gestor + admin", "Planilha de revisão assinada"],
    ["", "Trimestral", "Acessos da org GitHub, Azure RBAC, Key Vault e ACL do broker", "Segurança", "Export de acessos"],
    ["Rotação de segredos", "90 dias / anual", "JWT_SECRET e certificados IoT a cada 90 dias; DATA_KEY anual (recifra via ID de chave)",
     "DevOps", "Histórico de versões do Key Vault"],
    ["Backup e recuperação (continuação da rotina iniciada na Fase 3)", "Diário",
     "Backup automático do PostgreSQL + PITR de 7 dias, cópia geo-redundante cifrada, retenção de 35 dias", "DevOps", "Relatório do Azure Backup"],
    ["", "Mensal", "Teste de restore em ambiente isolado. Meta: RPO ≤ 24 h (PITR ~5 min), RTO ≤ 4 h", "DevOps", "Ata do teste com tempo medido"],
    ["", "Contínuo", "Key Vault com soft-delete e purge protection (sem a chave, o backup cifrado se perde)", "Segurança", "Configuração do Key Vault"],
], [3.0, 2.4, 6.6, 2.2, 2.8], font=7.5)

h2("4.8 Checklist de conformidade")
p("Legenda: **✔ Verificado** (implementado e executado/testado nesta Sprint) · **◐ Configurado** (arquivo "
  "pronto; roda quando o repositório for publicado e a infraestrutura provisionada) · **○ Planejado**.", size=9.5)
checklist = [
    ["Pipeline", "Secret scanning em todo o histórico, com regras próprias", "✔ / ◐", "1.5 (execução local) · devsecops.yml"],
    ["Pipeline", "SAST com regras próprias testadas", "✔ / ◐", "1.5 · .semgrep/"],
    ["Pipeline", "SCA + Dependabot", "✔ / ◐", "pip-audit, Trivy fs · dependabot.yml"],
    ["Pipeline", "Container + IaC scan, SBOM, assinatura", "✔ / ◐", "Trivy config executado · devsecops.yml (etapa 5)"],
    ["Pipeline", "Actions fixadas por SHA, OIDC, permissões mínimas", "◐", "devsecops.yml"],
    ["Pipeline", "DAST em staging e aprovação para produção", "◐", "devsecops.yml (etapas 7-8)"],
    ["Código", "Criptografia AES-256-GCM de PII + rotação de chave", "✔", "api/app/crypto.py · TestDataProtection"],
    ["Código", "JWT seguro (alg fixo, claims, expiração, revogação)", "✔", "api/app/security.py · TestAuthentication"],
    ["Código", "Rate limit, limite de corpo e validação por allow-list", "✔", "api/app/ · TestHardening"],
    ["Código", "RBAC com 3 perfis + proteção BOLA", "✔", "api/app/security.py · TestAuthorization"],
    ["Código", "Erros padronizados sem vazamento", "✔", "api/app/errors.py · TestHardening"],
    ["ML", "Integridade do modelo (hash obrigatório em prod)", "✔", "api/app/model.py, config.py · TestDataProtection/TestHardening"],
    ["Mobile", "Token em armazenamento seguro, só HTTPS, versão mínima", "✔ (API) / ◐ (app)", "secureSession.ts · teste 426"],
    ["Mobile", "Certificate pinning e detecção de root", "○", "Build nativo (EAS) - próxima fase"],
    ["IoT", "TLS 1.2+, mTLS, ACL e CRL", "✔ (mTLS) / ◐ (broker)", "iot/mosquitto/ · test_mtls"],
    ["IoT", "Ingestão com anti-replay, esquema e janela de tempo", "✔", "iot/ingestion.py · test_ingestion"],
    ["Infra", "Dockerfiles não-root, multi-stage, sem segredos", "✔", "Trivy config: 0 achados"],
    ["Infra", "K8s hardening + NetworkPolicy", "✔", "Trivy config: 0 achados (1 aceite com prazo)"],
    ["Observab.", "Logs JSON com auditoria e redação de PII", "✔", "api/app/logging_setup.py · TestAuditLogging"],
    ["Observab.", "Métricas + 11 alertas (API, mobile, IoT, ML)", "✔", "alert_rules.yml · avaliar_alertas.py"],
    ["Observab.", "Dashboard Grafana + Loki + Alertmanager", "◐", "observability/, infra/docker-compose.yml"],
    ["Observab.", "Plano de resposta com 6 playbooks", "✔", "Seção 3.4"],
    ["Compliance", "STRIDE + riscos DevSecOps + OWASP ASVS/API/Mobile", "✔", "Seções 4.1 a 4.5"],
    ["LGPD", "Minimização, pseudonimização, cifra, incidente", "✔", "Seção 4.6"],
    ["LGPD", "RIPD, inventário (ROPA), direitos do titular", "○", "Jurídico/DPO + backlog"],
    ["Contínuo", "Rotinas de dependências, testes, permissões e backup", "✔ (definido)", "Seção 4.7"],
]
table(["Área", "Item", "Status", "Evidência"], checklist, [1.9, 6.3, 2.6, 6.2], font=7.6)

# ================================================================== ANEXOS
h1("Anexo A. Estrutura do repositório e como executar")
code("""cyber/
├── .github/workflows/devsecops.yml   pipeline DevSecOps (8 etapas, Actions fixadas por SHA)
├── .github/dependabot.yml            SCA contínuo (pip, npm, Docker, compose, Actions)
├── .gitleaks.toml  .semgrep/  .zap/  .trivyignore   regras e exceções das ferramentas
├── api/                              API FastAPI (app/, tests/, Dockerfile)
├── mobile/                           sessão segura + hardening Android (Expo)
├── iot/                              broker (TLS/mTLS/ACL/CRL), ingestão, cliente TLS, PKI, tests/
├── infra/                            docker-compose (stack completa), k8s/api.yaml, .env.example
├── observability/                    Prometheus, Alertmanager, Loki, Alloy, Grafana, simulador, avaliador
└── docs/                             este documento, figuras, evidências e geradores""", size=7.6)
code("""# 1) Ambiente e testes
python -m venv .venv
.venv\\Scripts\\pip install -r api\\requirements-dev.txt -r iot\\requirements.txt
cd api; ..\\.venv\\Scripts\\python -m pytest -v; cd ..\\iot; ..\\.venv\\Scripts\\python -m pytest -v; cd ..

# 2) Ferramentas de segurança (as mesmas do pipeline)
.venv\\Scripts\\pip install semgrep bandit pip-audit
.venv\\Scripts\\python .semgrep\\verificar_regras.py
.venv\\Scripts\\semgrep scan --config .semgrep/vinshare-rules.yml --config p/owasp-top-ten --exclude .semgrep api iot mobile
.venv\\Scripts\\bandit -r api/app iot -x iot/tests -ll -ii
.venv\\Scripts\\pip-audit -r api\\requirements.txt -r iot\\requirements.txt

# 3) Simulação de ataques, validação dos alertas, figuras e documento
.venv\\Scripts\\python observability\\simular_trafego.py
.venv\\Scripts\\python observability\\avaliar_alertas.py
.venv\\Scripts\\python docs\\gerar_figuras.py; .venv\\Scripts\\python docs\\gerar_documento.py""",
     "Comandos (PowerShell, a partir da pasta cyber/)", size=7.4)

h1("Anexo B. Lista completa dos testes automatizados")
table(["#", "Categoria", "Teste", "Resultado"],
      [[str(i + 1), labels.get(c, (c,))[0], name, r] for i, (c, name, r) in enumerate(all_cases)],
      [0.9, 4.6, 9.3, 2.2], font=7.0)

doc.save(OUT)
print("Documento gerado:", OUT)
